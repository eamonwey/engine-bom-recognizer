"""Regressions through the real pipeline, save/export routes and review store."""
import copy
import csv
import hashlib
import io
import json
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image
from engine_bom import webapp as w
from engine_bom.agent_loop import AgentLoop, validate_bom_response
from engine_bom.config import DEFAULT_CONFIG, load_config
from engine_bom.docextract import EXTRACT_SCHEMA, _extract_chunked
from engine_bom.export import write_json
from engine_bom.knowledge import build_knowledge
from engine_bom.pipeline import BomPipeline
from engine_bom.providers import ProviderResponse
from engine_bom.review_workflow import ReviewStore, digest
from engine_bom.schema import BomComponent, BomDocument, check_standard_references, component_from_dict, validate_components
from engine_bom.inspect_rules import check_bom_table


class FixedProvider:
    name = 'test'

    def __init__(self, row):
        self.raw = json.dumps({'components': [row]}, ensure_ascii=False)
        self.calls = []

    def analyze(self, images, prompt):
        self.calls.append(prompt)
        return ProviderResponse(self.raw, 'test')

    def describe(self):
        return {'provider': self.name}


class ExtractionRegressions(unittest.TestCase):
    def test_review_findings_survive_complete_pipeline_without_retry_or_rewrite(self):
        rows = [{'item_no': 1, 'name_zh': '轴'},
                {'item_no': 1, 'part_number': 'GB/T 5780-2016', 'name_zh': '轴承6206', 'quantity': 1}]
        with tempfile.TemporaryDirectory() as directory:
            source = os.path.join(directory, 'source.png')
            Image.new('RGB', (32, 32), 'white').save(source)
            for row in rows:
                with self.subTest(row=row):
                    cfg = load_config()
                    cfg['knowledge']['enabled'] = False
                    cfg['assembly']['enabled'] = False
                    cfg['output'].update(dir=directory, formats=['json'])
                    pipeline = BomPipeline(cfg)
                    provider = FixedProvider(row)
                    pipeline.provider = provider
                    pipeline.agent = AgentLoop(provider)
                    result = pipeline.process_file(source)
                    self.assertEqual(result.status, 'ok', result.errors)
                    self.assertEqual(len(provider.calls), 1)
                    component = result.bom.components[0]
                    self.assertTrue(component.needs_review)
                    self.assertEqual(component.part_number, row.get('part_number', ''))
                    persisted = json.loads(Path(directory, 'source.json').read_text(encoding='utf-8'))
                    self.assertTrue(persisted['components'][0]['review_reasons'])

    def test_review_filter_does_not_accept_unidentified_rows(self):
        self.assertTrue(validate_bom_response('{"components":[{}]}')[1])

    def test_year_and_width_variants_preserve_reference_and_find_mismatch(self):
        for reference in ('GB/T 5780-2016', 'GB / T 5780—2016', 'ＧＢ／Ｔ ５７８０－２０１６'):
            with self.subTest(reference=reference):
                row = BomComponent(part_number=reference, name_zh='轴承6206', quantity=1)
                self.assertEqual(len(check_standard_references([row])), 1)
                self.assertEqual(row.part_number, reference)
                row.name_zh = '螺栓M8'
                self.assertEqual(check_standard_references([row]), [])

    def test_explicit_quantity_markers_agree_with_inspection(self):
        for marker in ('AR', 'A/R', 'AS REQUIRED', '按需', '见注'):
            row = BomComponent(item_no=1, name_zh='润滑油', remark='数量原文: ' + marker)
            with self.subTest(marker=marker):
                self.assertEqual(validate_components([row]), [])
                self.assertEqual(check_bom_table([row.to_dict()]), [])
        row.remark = ''
        self.assertTrue(validate_components([row]))

    def test_nonfinite_model_weight_requires_retry_and_cannot_export_nan(self):
        for weight in ('NaN', 'Infinity', '-Infinity', float('nan')):
            row, warnings = component_from_dict({'part_number': 'A', 'weight': weight})
            self.assertIsNone(row.weight)
            self.assertTrue(warnings)
            self.assertTrue(validate_bom_response(json.dumps({'components': [{'part_number': 'A', 'weight': weight}]}))[1])
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(ValueError):
                write_json(BomDocument('a', components=[BomComponent(part_number='A', weight=float('nan'))]), os.path.join(directory, 'nan.json'))

    def test_long_block_sources_have_distinct_slices_and_offsets(self):
        class Provider:
            name = 'test'

            def analyze(self, images, prompt):
                batch = json.loads(prompt.split('来源块：', 1)[1])
                value = {key: [] for key in EXTRACT_SCHEMA}
                value['tech_params'] = [{'name': '参数', 'source_block_id': batch[0]['id']}]
                return ProviderResponse(json.dumps(value), 'test')

        text = 'A' * 9000 + 'B' * 200
        result = _extract_chunked(Provider(), {'blocks': [{'id': 'b1', 'text': text, 'source_page': 1}]}, [], {})
        rows = result['tech_params']
        self.assertEqual(len({row['source_block_id'] for row in rows}), 2)
        self.assertEqual(''.join(row['source_text'] for row in rows), text)
        self.assertEqual([(r['source_start'], r['source_end']) for r in rows], [(0, 9000), (9000, 9200)])
        self.assertTrue(all(r['source_parent_block_id'] == 'b1' for r in rows))


class WebRegressions(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        row = BomComponent(item_no=1, part_number='A', name_zh='轴', quantity=0)
        findings = validate_components([row])
        self.job = {'kind': 'bom', 'status': 'done', 'dir': str(self.root), 'result': {
            'document_id': 'a', 'bom': {'components': [row.to_dict()], 'model_info': {'preserve_printed_rows': True},
            'validation_issues': findings + ['标题栏仍待核实']}}}
        for key, value in [('DATA_DIR', str(self.root)), ('REVIEWLOG_PATH', str(self.root / 'log.json')), ('_jobs', {'j': self.job})]:
            p = patch.object(w, key, value)
            p.start()
            self.addCleanup(p.stop)
        self.handler = object.__new__(w.Handler)
        self.responses = []
        self.handler._json = lambda value, code=200: self.responses.append((code, value))

    def test_save_refreshes_flags_and_persisted_findings_in_both_directions(self):
        for quantity in (2, 0, 3):
            rows = copy.deepcopy(self.job['result']['bom']['components'])
            rows[0]['quantity'] = quantity
            self.handler._read_json_body = lambda: {'job_id': 'j', 'document_id': 'a', 'components': rows}
            self.handler._api_bom_save()
            bom = self.job['result']['bom']
            self.assertEqual(bom['components'][0]['needs_review'], quantity == 0)
            self.assertEqual(bool(bom['components'][0]['review_reasons']), quantity == 0)
            quantity_notes = [s for s in bom['validation_issues'] if '数量非正数' in s]
            self.assertEqual(len(quantity_notes), int(quantity == 0))
            self.assertIn('标题栏仍待核实', bom['validation_issues'])
            self.assertEqual(json.loads((self.root / 'result.json').read_text(encoding='utf-8')), self.job['result'])
            saved = json.loads((self.root / 'reviewed' / 'a.json').read_text(encoding='utf-8'))
            self.assertEqual(saved['components'], bom['components'])

    def test_enterprise_exports_handle_empty_and_nonempty_reason_lists(self):
        from openpyxl import load_workbook
        self.handler._job_result = lambda _: self.job['result']
        output = []
        self.handler._send = lambda code, body, *args: output.append((code, body))
        for reasons in ([], ['数量非正数', '标准号待核实']):
            self.job['result']['bom']['components'][0].update(review_reasons=reasons, needs_review=bool(reasons))
            for fmt in ('csv', 'xlsx'):
                self.handler.path = '/api/export/bom/j?profile=enterprise&fmt=' + fmt
                with patch.object(w, 'load_web_config', return_value={'enterprise': {}}):
                    self.handler._api_export_bom('j')
                code, body = output[-1]
                self.assertEqual(code, 200)
                if fmt == 'csv':
                    rows = list(csv.reader(io.StringIO(body.decode('utf-8-sig'))))
                else:
                    book = load_workbook(io.BytesIO(body))
                    rows = list(book.active.values)
                    book.close()
                reason = rows[1][list(rows[0]).index('核实原因')]
                self.assertEqual(reason or '', '；'.join(reasons))

    def test_legacy_ar_flags_are_cleared_without_invalidating_unchanged_relations(self):
        old_note = '#1 润滑油: 数量、材料、件号均为空——疑似无明细表页的占位行，请对照原图核实'
        row = BomComponent(item_no=1, name_zh='润滑油', remark='数量原文: AR', needs_review=True, review_reasons=[old_note])
        bom = self.job['result']['bom']
        bom.update(components=[row.to_dict()], validation_issues=[old_note, '标题栏仍待核实'], revision=4,
                   assembly={'nodes': [{'id': 'n', 'parent_id': None, 'component_index': 0}]})
        self.handler._read_json_body = lambda: {'job_id': 'j', 'document_id': 'a', 'components': copy.deepcopy(bom['components'])}
        self.handler._api_bom_save()
        updated = self.job['result']['bom']
        self.assertFalse(updated['components'][0]['needs_review'])
        self.assertNotIn(old_note, updated['validation_issues'])
        self.assertEqual(updated['assembly']['nodes'], bom['assembly']['nodes'])
        self.assertEqual(updated['revision'], 4)

    def test_balloon_bad_numbers_are_review_notes_not_matches_or_crashes(self):
        for value in (2.9, float('nan'), float('inf'), float('-inf')):
            with self.subTest(value=value):
                self.assertIn('非法数值', w.balloon_bom_mismatch_note([value], [{'item_no': 2}]))
        self.assertIsNone(w.balloon_bom_mismatch_note([2.0, '3'], [{'item_no': 2}, {'item_no': 3}]))

    def test_missing_config_does_not_share_nested_defaults(self):
        original = copy.deepcopy(DEFAULT_CONFIG)
        with patch.object(w, '_ROOT', str(self.root)), patch.object(w, 'CONFIG_PATH', None), patch.object(w, 'SETTINGS_PATH', str(self.root / 'missing.json')):
            first = w.load_web_config()
            first['provider']['name'] = 'first_job'
            first['ingestion']['preprocess']['enabled'] = True
            self.assertEqual(w.load_web_config(), original)
        self.assertEqual(DEFAULT_CONFIG, original)

    def test_standalone_inspection_anchors_kb_and_passes_common_options(self):
        cfg = load_config()
        cfg['enterprise']['required_annotations'] = [{'id': 'X'}]
        knowledge = self.root / 'rules.txt'
        knowledge.write_text('轴承规则\n', encoding='utf-8')
        cfg['knowledge']['path'] = 'rules.txt'
        captured = []
        def pipeline(config, base_dir):
            captured.append(build_knowledge(config['knowledge'], base_dir).describe())
            return SimpleNamespace(process_file=lambda _: SimpleNamespace(status='ok', bom=SimpleNamespace(to_dict=lambda: {'components': []})))
        inspect_result = {'issues': [], 'visual_status': 'validated'}
        source = self.root / 'drawing.png'
        Image.new('RGB', (8, 8)).save(source)
        with patch.object(w, '_ROOT', str(self.root)), patch.object(w, 'load_web_config', return_value=cfg), \
             patch('engine_bom.ingestion.load_document', return_value=[SimpleNamespace(image=Image.new('RGB', (8, 8)))]), \
             patch('engine_bom.providers.build_provider', return_value=SimpleNamespace(describe=lambda: {})), \
             patch('engine_bom.titleblock.extract_titleblock', return_value={'fields': {}, 'missing_required': []}), \
             patch('engine_bom.titleblock.extract_dims_and_balloons', return_value={}), \
             patch('engine_bom.pipeline.BomPipeline', side_effect=pipeline), \
             patch('engine_bom.hybrid_inspection.inspect', return_value=inspect_result) as inspect:
            w._job_inspect(str(self.root / 'job'), str(source), {})
        self.assertEqual(captured[0]['entries'], 1)
        self.assertEqual(inspect.call_args[0][-1]['required_annotations'], [{'id': 'X'}])
        self.assertTrue(inspect.call_args[0][-1]['check_balloons'])

    def test_review_get_requires_api_and_review_identity_tokens(self):
        store = ReviewStore(str(self.root / 'reviews.sqlite'))
        store.transition('a', 'submit', {'id': 'e', 'role': 'editor'}, payload={'quantity': 2})
        token = 'test-reviewer'
        users = [{'id': 'r', 'role': 'reviewer', 'token_sha256': hashlib.sha256(token.encode()).hexdigest()}]
        (self.root / 'review_users.json').write_text(json.dumps(users), encoding='utf-8')
        self.handler.path = '/api/reviews/a'
        for headers, status in [({}, 403), ({'X-API-Token': w._API_TOKEN}, 403),
                                ({'X-API-Token': w._API_TOKEN, 'X-Review-Token': 'wrong'}, 403),
                                ({'X-API-Token': w._API_TOKEN, 'X-Review-Token': token}, 200)]:
            self.handler.headers = headers
            self.handler.do_GET()
            self.assertEqual(self.responses[-1][0], status)


class ReviewDigestRegressions(unittest.TestCase):
    def test_equivalent_numbers_allow_review_and_approval_but_changes_do_not(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ReviewStore(os.path.join(directory, 'reviews.sqlite'))
            editor, reviewer, approver = [dict(id=role, role=role) for role in ('editor', 'reviewer', 'approver')]
            original = {'components': [{'weight': 2.0, 'quantity': 1}], 'enabled': True}
            browser = {'components': [{'weight': 2, 'quantity': 1.0}], 'enabled': True}
            store.transition('a', 'submit', editor, payload=original)
            self.assertEqual(store.transition('a', 'review', reviewer, 1, browser)['state'], 'reviewed')
            changed = copy.deepcopy(browser)
            changed['components'][0]['weight'] = 2.5
            with self.assertRaises(ValueError):
                store.transition('a', 'approve', approver, 2, changed)
            self.assertEqual(store.transition('a', 'approve', approver, 2, browser)['state'], 'approved')
            self.assertNotEqual(digest({'x': True}), digest({'x': 1}))

    def test_legacy_float_snapshot_remains_reviewable_without_rewriting_audit_hash(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, 'reviews.sqlite')
            store = ReviewStore(path)
            payload = {'weight': 2.0}
            store.transition('a', 'submit', {'id': 'e', 'role': 'editor'}, payload=payload)
            legacy = hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True, allow_nan=False).encode('utf-8')).hexdigest()
            with sqlite3.connect(path) as db:
                for table in ('versions', 'snapshots', 'audit'):
                    db.execute('UPDATE %s SET hash=?' % table, (legacy,))
            db.close()
            result = store.transition('a', 'review', {'id': 'r', 'role': 'reviewer'}, 1, {'weight': 2})
            self.assertEqual(result['state'], 'reviewed')
            self.assertEqual(result['hash'], legacy)
            self.assertTrue(all(row['hash'] == legacy for row in result['audit']))


if __name__ == '__main__':
    unittest.main()
